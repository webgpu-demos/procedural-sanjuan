#!/usr/bin/env python3
"""The Census data the demand model (demand.py) is built from, downloaded and cut to Puerto Rico:
  pr_flows.json      ACS 2016-2020 commuting flows, residence municipio -> work municipio (table 1)
  departures.json    ACS 2018-2022 5-year B08302, time leaving home to go to work, by municipio
  municipios.json    the 78 municipios' boundaries (cartographic boundary file 2022, 1:500k)
All public domain (U.S. Census Bureau). Writes data/sim/census/. Usage: .venv-sim/bin/python tools/sim/census.py
"""
import io, json, os, urllib.request, zipfile
import openpyxl, shapefile

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), '../..'))
OUT = os.path.join(ROOT, 'data/sim/census'); os.makedirs(OUT, exist_ok=True)
get = lambda url: urllib.request.urlopen(url, timeout=600).read()

# ---- commuting flows
wb = openpyxl.load_workbook(io.BytesIO(get('https://www2.census.gov/programs-surveys/demo/tables/metro-micro/2020/commuting-flows-2020/table1.xlsx')), read_only=True)
flows = []
for r in wb[wb.sheetnames[0]].iter_rows(values_only=True):
    if r and str(r[0]) == '72' and r[5] is not None:
        pr = r[6] == 'Puerto Rico'
        flows.append({'from': r[3].replace(' Municipio', ''), 'to': r[7].replace(' Municipio', '') if pr else (r[6] or '?'), 'n': r[8], 'moe': r[9],
                      'toFips': str(r[5]) if pr else None, 'fromFips': str(r[1])})
json.dump(flows, open(os.path.join(OUT, 'pr_flows.json'), 'w'))
print(len(flows), 'flows from Puerto Rico municipios')

# ---- boundaries (first: the departures are keyed by name)
z = zipfile.ZipFile(io.BytesIO(get('https://www2.census.gov/geo/tiger/GENZ2022/shp/cb_2022_us_county_500k.zip')))
sf = shapefile.Reader(shp=io.BytesIO(z.read('cb_2022_us_county_500k.shp')), dbf=io.BytesIO(z.read('cb_2022_us_county_500k.dbf')))
fields = [f[0] for f in sf.fields[1:]]
munis, names = {}, {}
for sr in sf.iterShapeRecords():
    r = dict(zip(fields, sr.record))
    if r['STATEFP'] != '72': continue
    pts, parts = sr.shape.points, list(sr.shape.parts) + [len(sr.shape.points)]
    munis[r['NAME']] = {'fips': r['COUNTYFP'], 'rings': [[[round(x, 5), round(y, 5)] for x, y in pts[parts[i]:parts[i + 1]]] for i in range(len(parts) - 1)]}
    names[r['COUNTYFP']] = r['NAME']
json.dump(munis, open(os.path.join(OUT, 'municipios.json'), 'w'))
print(len(munis), 'municipios')

# ---- time leaving home (the table-based summary file: every geography; keep Puerto Rico's municipios)
LABELS = ['total', '00-05', '05-0530', '0530-06', '06-0630', '0630-07', '07-0730', '0730-08', '08-0830', '0830-09', '09-10', '10-11', '11-12', '12-16', '16-24']
deps = {}
for line in get('https://www2.census.gov/programs-surveys/acs/summary_file/2022/table-based-SF/data/5YRData/acsdt5y2022-b08302.dat').decode().splitlines():
    if not line.startswith('0500000US72'): continue
    p = line.split('|'); fips = p[0][-3:]
    if fips in names: deps[names[fips]] = dict(zip(LABELS, [int(p[1 + 2 * i]) for i in range(15)]))
json.dump(deps, open(os.path.join(OUT, 'departures.json'), 'w'), indent=1)
print(len(deps), 'municipios with departure times')
